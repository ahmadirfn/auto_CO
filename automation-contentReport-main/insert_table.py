
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from PIL import Image

import subprocess
import tempfile
import shutil


# ============================================================
# 1. PATH CONFIGURATION
# ============================================================

BASE_DIR = Path.cwd()

INPUT_DIR = BASE_DIR / "input"
ASSETS_DIR = INPUT_DIR / "assets"

PPTX_PATH = INPUT_DIR / "master_template.pptx"
OUTPUT_PATH = BASE_DIR / "output.pptx"

IMAGE_PATH = (
    ASSETS_DIR
    / "KVs"
    / "01 Front AMG Line"
    / "12344158_C.P.BH.X540_PAD_BEV_ProdStillUnreal_01-EXT-3-4-Front-AMG-Line_1-1_plate-DE_RGB preview.jpg"
)

VIDEO_PATH = (
    ASSETS_DIR
    / "Film_IAA"
    / "12344422_C.P.BH.X540_PAD_BEV_ProdFilm_07-INT-IAA-Trimpart-10s_1-1_Clean-NOSound_ProRes.mp4"
)


# ============================================================
# 2. SLIDE AND SHAPE CONFIGURATION
# ============================================================

# Slide 2 - Image
IMAGE_SLIDE_INDEX = 1
IMAGE_SHAPE_ID = 10

# Slide 3 - Table
TABLE_SLIDE_INDEX = 2
TABLE_SHAPE_ID = 8

# Slide 5 - Video
VIDEO_SLIDE_INDEX = 4
VIDEO_SHAPE_ID = 6

# Video thumbnail timestamp in seconds
THUMBNAIL_TIMESTAMP = 2


# ============================================================
# 3. TABLE DATA
# ============================================================

TABLE_DATA = [
    ["Asset ID", "Variation", "Quality", "Symphony Link"],
    ["12344158", "Front AMG Line", "Approved", "Link 1"],
    ["12344422", "IAA Trimpart", "Approved", "Link 2"]
]


# ============================================================
# 4. VALIDATE INPUT FILES
# ============================================================

def validate_files():

    for path in [PPTX_PATH, IMAGE_PATH, VIDEO_PATH]:
        if not path.is_file():
            raise FileNotFoundError(
                f"File not found: {path}"
            )

    if shutil.which("ffmpeg") is None:
        raise RuntimeError(
            "FFmpeg not found. Install FFmpeg and add it to PATH."
        )

    print("All input files validated.")


# ============================================================
# 5. FIND SHAPE BY ID
# ============================================================

def find_shape(slide, shape_id):

    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape

    raise ValueError(
        f"Shape ID {shape_id} not found."
    )


# ============================================================
# 6. REPLACE ORIGINAL SHAPE
# ============================================================

def replace_shape(original_shape, new_shape):

    old_element = original_shape._element
    new_element = new_shape._element

    parent = old_element.getparent()
    index = parent.index(old_element)

    # Remove new shape from current position
    new_parent = new_element.getparent()
    new_parent.remove(new_element)

    # Insert at original stacking position
    parent.insert(index, new_element)

    # Remove original placeholder
    parent.remove(old_element)


# ============================================================
# 7. INSERT IMAGE
# ============================================================

def insert_image(slide, shape_id, image_path):

    shape = find_shape(slide, shape_id)

    print(f"\nImage placeholder found: {shape.name}")

    if (
        shape.is_placeholder
        and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ):

        shape.insert_picture(str(image_path))

        print("Image inserted into native picture placeholder.")
        return

    picture = slide.shapes.add_picture(
        str(image_path),
        shape.left,
        shape.top,
        width=shape.width,
        height=shape.height
    )

    replace_shape(shape, picture)

    print("Image inserted successfully.")


# ============================================================
# 8. INSERT TABLE
# ============================================================

def insert_table(slide, shape_id, data):

    if not data or not data[0]:
        raise ValueError("Table data cannot be empty.")

    rows = len(data)
    cols = len(data[0])

    if not all(len(row) == cols for row in data):
        raise ValueError(
            "All table rows must have the same number of columns."
        )

    shape = find_shape(slide, shape_id)

    print(f"\nTable placeholder found: {shape.name}")

    # Native table placeholder
    if (
        shape.is_placeholder
        and shape.placeholder_format.type == PP_PLACEHOLDER.TABLE
    ):

        graphic_frame = shape.insert_table(
            rows=rows,
            cols=cols
        )

        table = graphic_frame.table

        print("Table inserted into native table placeholder.")

    else:

        graphic_frame = slide.shapes.add_table(
            rows=rows,
            cols=cols,
            left=shape.left,
            top=shape.top,
            width=shape.width,
            height=shape.height
        )

        table = graphic_frame.table

        replace_shape(shape, graphic_frame)

        print("Table inserted by replacing placeholder.")

    # --------------------------------------------------------
    # POPULATE TABLE
    # --------------------------------------------------------

    for row_idx, row_data in enumerate(data):

        for col_idx, value in enumerate(row_data):

            cell = table.cell(row_idx, col_idx)

            cell.text = (
                str(value) if value is not None else ""
            )

    # --------------------------------------------------------
    # FORMAT TABLE
    # --------------------------------------------------------

    for row_idx, row in enumerate(table.rows):

        for cell in row.cells:

            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            # Cell margins
            cell.margin_left = Pt(5)
            cell.margin_right = Pt(5)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)

            for paragraph in cell.text_frame.paragraphs:

                paragraph.alignment = PP_ALIGN.CENTER

                for run in paragraph.runs:

                    run.font.size = Pt(10)

                    if row_idx == 0:
                        run.font.bold = True

    # --------------------------------------------------------
    # SET COLUMN WIDTHS
    # --------------------------------------------------------

    if cols == 4:

        column_ratios = [0.20, 0.35, 0.15, 0.30]

        total_width = sum(
            col.width for col in table.columns
        )

        widths = [
            int(total_width * ratio)
            for ratio in column_ratios[:-1]
        ]

        widths.append(total_width - sum(widths))

        for col, width in zip(table.columns, widths):
            col.width = width

    print(
        f"Table inserted successfully: {rows} rows x {cols} columns"
    )

    return table


# ============================================================
# 9. EXTRACT VIDEO THUMBNAIL USING FFMPEG
# ============================================================

def extract_video_thumbnail(
    video_path,
    output_path,
    timestamp=2
):

    print(
        f"\nExtracting video thumbnail at {timestamp}s..."
    )

    command = [
        "ffmpeg",
        "-y",
        "-ss", str(timestamp),
        "-i", str(video_path),
        "-frames:v", "1",
        "-q:v", "2",
        str(output_path)
    ]

    result = subprocess.run(
        command,
        capture_output=True,
        text=True
    )

    if result.returncode != 0:
        raise RuntimeError(
            f"FFmpeg failed:\n{result.stderr}"
        )

    if not output_path.is_file():
        raise RuntimeError(
            "Thumbnail was not generated."
        )

    with Image.open(output_path) as img:
        img.verify()

    print(
        f"Thumbnail generated: {output_path}"
    )

    return output_path


# ============================================================
# 10. EMBED VIDEO WITH THUMBNAIL
# ============================================================

def insert_video(
    slide,
    shape_id,
    video_path,
    thumbnail_path
):

    shape = find_shape(slide, shape_id)

    print(
        f"\nVideo placeholder found: {shape.name}"
    )

    print(
        f"Position: ({shape.left}, {shape.top})"
    )

    print(
        f"Size: ({shape.width}, {shape.height})"
    )

    movie = slide.shapes.add_movie(
        movie_file=str(video_path),
        left=shape.left,
        top=shape.top,
        width=shape.width,
        height=shape.height,
        poster_frame_image=str(thumbnail_path),
        mime_type="video/mp4"
    )

    replace_shape(shape, movie)

    print(
        "Video embedded with thumbnail successfully."
    )


# ============================================================
# 11. MAIN AUTOMATION
# ============================================================

def main():

    print("=" * 60)
    print("CONTENT REPORT AUTOMATION")
    print("=" * 60)

    validate_files()

    # Open existing presentation
    prs = Presentation(str(PPTX_PATH))

    print(
        f"\nTotal slides: {len(prs.slides)}"
    )

    required_indices = [
        IMAGE_SLIDE_INDEX,
        TABLE_SLIDE_INDEX,
        VIDEO_SLIDE_INDEX
    ]

    if max(required_indices) >= len(prs.slides):
        raise IndexError(
            "A configured slide index exceeds the number of slides."
        )

    # --------------------------------------------------------
    # SLIDE 2 - INSERT IMAGE
    # --------------------------------------------------------

    print("\n" + "-" * 60)
    print("PROCESSING SLIDE 2 - IMAGE")
    print("-" * 60)

    image_slide = prs.slides[IMAGE_SLIDE_INDEX]

    insert_image(
        slide=image_slide,
        shape_id=IMAGE_SHAPE_ID,
        image_path=IMAGE_PATH
    )

    # --------------------------------------------------------
    # SLIDE 3 - INSERT TABLE
    # --------------------------------------------------------

    print("\n" + "-" * 60)
    print("PROCESSING SLIDE 3 - TABLE")
    print("-" * 60)

    table_slide = prs.slides[TABLE_SLIDE_INDEX]

    insert_table(
        slide=table_slide,
        shape_id=TABLE_SHAPE_ID,
        data=TABLE_DATA
    )

    # --------------------------------------------------------
    # SLIDE 5 - INSERT VIDEO
    # --------------------------------------------------------

    print("\n" + "-" * 60)
    print("PROCESSING SLIDE 5 - VIDEO")
    print("-" * 60)

    video_slide = prs.slides[VIDEO_SLIDE_INDEX]

    with tempfile.TemporaryDirectory() as temp_dir:

        thumbnail_path = (
            Path(temp_dir) / "video_thumbnail.jpg"
        )

        extract_video_thumbnail(
            video_path=VIDEO_PATH,
            output_path=thumbnail_path,
            timestamp=THUMBNAIL_TIMESTAMP
        )

        insert_video(
            slide=video_slide,
            shape_id=VIDEO_SHAPE_ID,
            video_path=VIDEO_PATH,
            thumbnail_path=thumbnail_path
        )

        # Save presentation while thumbnail is available
        prs.save(str(OUTPUT_PATH))

    # --------------------------------------------------------
    # COMPLETION
    # --------------------------------------------------------

    print("\n" + "=" * 60)
    print("AUTOMATION COMPLETED SUCCESSFULLY")
    print("=" * 60)

    print(f"Output file: {OUTPUT_PATH}")

    print(
        f"File size: "
        f"{OUTPUT_PATH.stat().st_size / (1024 * 1024):.2f} MB"
    )


# ============================================================
# 12. EXECUTE
# ============================================================

if __name__ == "__main__":
    main()
