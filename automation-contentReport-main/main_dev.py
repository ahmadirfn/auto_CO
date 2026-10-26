import pandas as pd
import numpy as np
from pptx import Presentation
import os

input_path = os.path.join(os.getcwd(), "input")
pptx_path = os.path.join(input_path, "master_template.pptx")

prs = Presentation(pptx_path)

for slide_no, slide in enumerate(prs.slides, start=1):
    print(f"\n===== SLIDE {slide_no} =====")

    for i, shape in enumerate(slide.shapes, start=1):
        print(f"\nObject {i}")
        print(f"  Name: {shape.name}")
        print(f"  ID: {shape.shape_id}")
        print(f"  Type: {shape.shape_type}")
        print(f"  Position: ({shape.left}, {shape.top})")
        print(f"  Size: ({shape.width}, {shape.height})")

        if shape.has_text_frame:
            print(f"  Text: {shape.text}")

        if shape.has_table:
            print("  Contains a table")