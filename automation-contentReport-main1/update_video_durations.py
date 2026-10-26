"""Fill missing video durations in Content_Overview_Enriched.xlsx from local videos.

Requires: pip install openpyxl; FFprobe installed and available in PATH.
Run: python update_video_durations.py [--assets PATH] [--excel PATH]
"""
from collections import defaultdict
import argparse
from pathlib import Path
import re
import shutil
import subprocess
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ASSETS_DIR = PROJECT_ROOT / "input" / "Assets"
EXCEL_PATH = PROJECT_ROOT / "input" / "Content_Overview_Enriched.xlsx"

VIDEO_EXTENSIONS = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".wmv", ".webm", ".mpg", ".mpeg"}


def missing(value):
    return value is None or isinstance(value, str) and not value.strip()


def normalized_name(name):
    return Path(str(name).strip().replace('\\', '/')).name.casefold()


def normalized_asset_id(value):
    if missing(value):
        return ""
    text = str(value).strip()
    # Excel may store an all-numeric Asset ID as a number (for example, 12345.0).
    if re.fullmatch(r"\d+\.0+", text):
        text = text.partition('.')[0]
    return text.casefold()


def video_index(assets_dir):
    by_name = defaultdict(list)
    by_stem = defaultdict(list)
    by_asset_id = defaultdict(list)
    for path in assets_dir.rglob('*'):
        if not path.is_file() or path.suffix.casefold() not in VIDEO_EXTENSIONS:
            continue
        by_name[path.name.casefold()].append(path)
        by_stem[path.stem.casefold()].append(path)
        match = re.match(r"^(\d+)(?:[_\-.]|$)", path.name)
        if match:
            by_asset_id[match.group(1).casefold()].append(path)
    return by_name, by_stem, by_asset_id


def find_video_candidates(filename, asset_id, index):
    by_name, by_stem, by_asset_id = index
    name = normalized_name(filename) if not missing(filename) else ""
    if name:
        candidates = by_name.get(name, [])
        if candidates:
            return list(dict.fromkeys(candidates))
        stem = Path(name).stem.casefold()
        candidates = by_stem.get(stem, [])
        if candidates:
            return list(dict.fromkeys(candidates))
    key = normalized_asset_id(asset_id)
    return list(dict.fromkeys(by_asset_id.get(key, []))) if key else []


def duration_seconds(path):
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
        capture_output=True, text=True, timeout=30, check=False,
    )
    if result.returncode != 0:
        return None
    try:
        seconds = float(result.stdout.strip())
        return int(seconds + 0.5) if 0 < seconds < float('inf') else None
    except ValueError:
        return None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--assets', type=Path, default=ASSETS_DIR,
                        help=f'Video asset directory (default: {ASSETS_DIR})')
    parser.add_argument('--excel', type=Path, default=EXCEL_PATH,
                        help=f'Enriched workbook (default: {EXCEL_PATH})')
    args = parser.parse_args()
    assets_dir = args.assets.expanduser().resolve()
    excel_path = args.excel.expanduser().resolve()
    if not assets_dir.is_dir():
        raise SystemExit(f"Assets directory not found: {assets_dir}")
    if not excel_path.is_file():
        raise SystemExit(f"Excel file not found: {excel_path}")
    if not shutil.which("ffprobe"):
        raise SystemExit("FFprobe not found. Install FFmpeg and add its bin folder to PATH.")

    print("Indexing video files recursively...")
    index = video_index(assets_dir)

    wb = load_workbook(excel_path)
    if "Items" not in wb.sheetnames:
        raise SystemExit("Worksheet 'Items' not found.")
    ws = wb["Items"]
    columns = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}
    required = ["Asset ID", "Filename", "Format", "Duration (seconds)", "Film Duration"]
    absent = [name for name in required if name not in columns]
    if absent:
        raise SystemExit(f"Missing columns: {absent}")

    updated = already_filled = not_found = ambiguous = probe_failed = 0
    examples = []
    for row in range(2, ws.max_row + 1):
        existing_numeric = ws.cell(row, columns["Duration (seconds)"])
        existing_formatted = ws.cell(row, columns["Film Duration"])
        if not missing(existing_numeric.value) and not missing(existing_formatted.value):
            already_filled += 1
            continue

        filename = ws.cell(row, columns["Filename"]).value
        asset_id = ws.cell(row, columns["Asset ID"]).value
        extension = Path(normalized_name(filename)).suffix.casefold() if not missing(filename) else ""
        if not extension:
            file_format = ws.cell(row, columns["Format"]).value
            extension = "." + str(file_format).strip().lstrip('.').casefold() if not missing(file_format) else ""
        if extension not in VIDEO_EXTENSIONS:
            continue  # Never fill durations for image/document assets.

        candidates = find_video_candidates(filename, asset_id, index)
        if not candidates:
            not_found += 1
            continue
        if len(candidates) != 1:
            ambiguous += 1
            examples.append(f"Row {row}: multiple video matches for {filename or asset_id}")
            continue

        seconds = duration_seconds(candidates[0])
        if seconds is None:
            probe_failed += 1
            examples.append(f"Row {row}: FFprobe failed for {candidates[0]}")
            continue

        # Preserve existing populated values. Only fill blanks.
        if missing(existing_numeric.value):
            existing_numeric.value = seconds
        if missing(existing_formatted.value):
            existing_formatted.value = f"{seconds}s"
        updated += 1

        # Remove ONLY the duration-specific notes that are now resolved.
        if "Mapping Notes" in columns:
            cell = ws.cell(row, columns["Mapping Notes"])
            if isinstance(cell.value, str):
                pieces = [p.strip() for p in cell.value.split(';') if p.strip()]
                pieces = [p for p in pieces if not p.casefold().startswith('film duration:')]
                cell.value = '; '.join(pieces) or None
        # Other notes and Group Key may still require re-running the mapper.
        # Do not incorrectly mark a record 'Mapped' here.

    if updated:
        wb.save(excel_path)  # Overwrite same workbook, as requested.
    print(f"Videos indexed: {sum(map(len, index[0].values()))}")
    print(f"Rows updated: {updated}")
    print(f"Rows already containing both durations: {already_filled}")
    print(f"Unmatched candidates: {not_found}; ambiguous: {ambiguous}; probe failed: {probe_failed}")
    for example in examples[:15]:
        print("REVIEW:", example)
    print("Saved:", excel_path if updated else "No changes")


if __name__ == "__main__":
    main()
