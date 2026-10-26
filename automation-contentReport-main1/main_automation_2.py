"""Generate the Mercedes-Benz content report from local media and two Excel sources.

Project layout:
  project/input/master_template.pptx
  project/input/assets_table1_cleaned.xlsx
  project/input/Export_report_20260917_113237_Raw(1).xlsx
  project/input/Assets/** (images and videos)
  project/automation-contentReport-main1/main_automation.py  (or project/main_automation.py)

Run: python main_automation.py
Plan without local media: python main_automation.py --plan-only
Dependencies: pip install python-pptx pandas openpyxl pillow imageio-ffmpeg
"""
from __future__ import annotations

import argparse
import csv
import math
import re
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Cm, Pt

IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tif', '.tiff', '.webp'}
VIDEO_EXTS = {'.mp4', '.mov', '.m4v', '.avi', '.wmv', '.webm', '.mkv'}
IMAGE_SLOTS = {'9:16': 3, '16:9': 11, '4:3': 12, '4:5': 13}
TABLE_HEADERS = ['Asset ID', 'Color Space', 'Aspect Ratio', 'Format', 'Symphony Link']
TABLE_FONT = 'MB Corpo S Text Office Light'
TABLE_HEADER_LABELS = {
    'Asset ID', 'Color Space', 'Aspect Ratio', 'Format', 'Symphony Link',
    'Variation', 'Quality',
}
TABLE_LAYOUT_WARNINGS = set()


def _text_width_pt(value, font_size):
    """Estimate text width in points using conservative per-character metrics."""
    width_units = 0.0
    for char in str(value):
        if char in ' ilI.,:;!|':
            width_units += 0.28
        elif char in 'MW@%&':
            width_units += 0.88
        elif char.isupper() or char.isdigit():
            width_units += 0.62
        elif char.isspace():
            width_units += 0.30
        else:
            width_units += 0.52
    return width_units * font_size


def _wrapped_line_count(value, width_pt, font_size):
    """Estimate natural word wrapping, splitting only words wider than a cell."""
    usable_width = max(1.0, width_pt)
    lines = 1
    line_width = 0.0
    for word in str(value).split():
        word_width = _text_width_pt(word, font_size)
        if word_width > usable_width:
            if line_width:
                lines += 1
                line_width = 0.0
            lines += max(0, math.ceil(word_width / usable_width) - 1)
            line_width = word_width % usable_width
        elif line_width and line_width + _text_width_pt(' ', font_size) + word_width > usable_width:
            lines += 1
            line_width = word_width
        else:
            line_width += (_text_width_pt(' ', font_size) if line_width else 0) + word_width
    return max(1, lines)


def _table_column_widths(data, available_width, font_size, margins_pt, wrap_columns=()):
    """Fit content-aware columns into the placeholder without losing text."""
    column_count = len(data[0])
    margin_width = 2 * margins_pt
    safety_pt = 1.5
    available_pt = available_width / 12700
    min_widths = []
    desired_widths = []
    for ci in range(column_count):
        values = [str(row[ci]) for row in data]
        if ci in wrap_columns:
            min_text = max((_text_width_pt(token, font_size)
                            for value in values for token in value.split()), default=font_size)
            desired_text = max((_text_width_pt(value, font_size) for value in values), default=font_size)
            minimum = min_text + safety_pt
        else:
            required = max((_text_width_pt(value, font_size) for value in values), default=font_size)
            desired_text = required
            minimum = required + safety_pt
        min_widths.append(minimum + margin_width)
        desired_widths.append(max(font_size, desired_text))

    # If the original font size cannot fit, reduce it in half-point steps as a
    # last resort. Never silently squeeze or truncate identifiers/metadata.
    while sum(min_widths) > available_pt and font_size > 7:
        font_size -= 0.5
        min_widths = []
        desired_widths = []
        for ci in range(column_count):
            values = [str(row[ci]) for row in data]
            if ci in wrap_columns:
                min_text = max((_text_width_pt(token, font_size)
                                for value in values for token in value.split()), default=font_size)
                desired_text = max((_text_width_pt(value, font_size) for value in values), default=font_size)
                minimum = min_text + safety_pt
            else:
                required = max((_text_width_pt(value, font_size) for value in values), default=font_size)
                desired_text = required
                minimum = required + safety_pt
            min_widths.append(minimum + margin_width)
            desired_widths.append(max(font_size, desired_text))
    if sum(min_widths) > available_pt:
        details = ', '.join(f'{data[0][i]} needs about {min_widths[i]:.1f} pt'
                            for i in range(column_count))
        raise ValueError(
            f'Metadata table requires {sum(min_widths):.1f} pt but its placeholder '
            f'provides {available_pt:.1f} pt; content was not truncated. {details}'
        )

    remaining = available_pt - sum(min_widths)
    weights = [max(1.0, desired_widths[i] - min_widths[i] + margin_width)
               for i in range(column_count)]
    weight_total = sum(weights)
    widths_pt = [min_widths[i] + remaining * weights[i] / weight_total
                 for i in range(column_count)]
    widths_emu = [int(width * 12700) for width in widths_pt[:-1]]
    widths_emu.append(available_width - sum(widths_emu))
    return widths_emu, font_size


def project_root():
    here = Path(__file__).resolve().parent
    for folder in (here, here.parent):
        if (folder / 'input' / 'master_template.pptx').is_file():
            return folder
    raise FileNotFoundError('Could not find input/master_template.pptx beside or above this script')


def asset_id(filename):
    match = re.match(r'^(\d+)', str(filename).strip())
    return match.group(1) if match else ''


def parse_ratio(value):
    try:
        a, b = str(value).strip().split(':')
        return float(a) / float(b)
    except (ValueError, ZeroDivisionError):
        return None


def media_kind(row):
    fmt = str(row.get('format', '')).lower().strip().lstrip('.')
    ext = Path(str(row['file_name'])).suffix.lower()
    if fmt in ('image', 'picture') or '.' + fmt in IMAGE_EXTS or ext in IMAGE_EXTS:
        return 'image'
    if fmt == 'video' or '.' + fmt in VIDEO_EXTS or ext in VIDEO_EXTS:
        return 'video'
    return None


def read_assets(excel, raw):
    df = pd.read_excel(excel, dtype=str).fillna('')
    df.columns = df.columns.astype(str).str.strip().str.lower()
    required = {'file_name', 'format', 'aspect_ratio'}
    if not required.issubset(df.columns):
        raise ValueError(f'Missing columns: {sorted(required - set(df.columns))}')
    meta = pd.read_excel(raw, dtype={'Asset ID': str}).fillna('')
    meta.columns = meta.columns.astype(str).str.strip()
    if 'Asset ID' not in meta.columns:
        raise ValueError('Raw export requires Asset ID column')
    lookup = {str(r['Asset ID']).strip(): r for _, r in meta.iterrows()}
    records = []
    for i, r in df.iterrows():
        if not str(r['file_name']).strip():
            continue
        kind = media_kind(r)
        if not kind:
            print(f'SKIP unknown format, Excel row {i+2}: {r["file_name"]}')
            continue
        aid = asset_id(r['file_name'])
        info = lookup.get(aid, {})
        instrument = str(info.get('Instruments', '')).strip() or 'Unknown Instrument'
        records.append(dict(index=i+2, name=str(r['file_name']).strip(), id=aid,
                            kind=kind, ratio=str(r['aspect_ratio']).strip(),
                            instrument=instrument, raw=info, source=r))
    return records


def find_local_files(records, folder):
    files = sorted((p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in IMAGE_EXTS | VIDEO_EXTS), key=lambda p: str(p).lower())
    by_name = defaultdict(list)
    by_stem = defaultdict(list)
    by_id = defaultdict(list)
    for p in files:
        by_name[p.name.casefold()].append(p)
        by_stem[p.stem.casefold()].append(p)
        by_id[asset_id(p.name)].append(p)
    for r in records:
        n = r['name'].casefold()
        stem = Path(n).stem if Path(n).suffix.lower() in IMAGE_EXTS | VIDEO_EXTS else n
        candidates = by_name[n] or by_stem[stem]
        if not candidates:
            # The metadata export sometimes truncates long filenames.
            candidates = [p for p in by_id[r['id']] if p.name.casefold().startswith(stem)]
        candidates = [p for p in candidates if (p.suffix.lower() in IMAGE_EXTS) == (r['kind'] == 'image')]
        if len(candidates) == 1:
            r['path'] = candidates[0]
        else:
            r['path'] = None
            r['problem'] = 'missing' if not candidates else 'ambiguous: ' + '; '.join(str(p) for p in candidates[:5])
    return records


def group_and_allocate(records):
    # Group by full Instruments string. Within each group, reserve exact slots before fallbacks.
    groups = defaultdict(list)
    for r in records:
        if r.get('path'):
            groups[r['instrument']].append(r)
    plan = []
    for instrument, items in groups.items():
        non_square = [r for r in items if r['kind'] == 'image' and r['ratio'] != '1:1']
        squares = [r for r in items if r['kind'] == 'image' and r['ratio'] == '1:1']
        videos = [r for r in items if r['kind'] == 'video']
        image_pages = []
        # First allocate exact / near exact ratio slots, then unmatched assets.
        matched, fallback = [], []
        for r in non_square:
            value = parse_ratio(r['ratio'])
            exact = [s for s in IMAGE_SLOTS if value and abs(math.log(value / parse_ratio(s))) <= 0.035]
            (matched if exact else fallback).append((r, exact[0] if exact else None))
        for r, preferred in matched + fallback:
            candidates = [preferred] if preferred else sorted(IMAGE_SLOTS, key=lambda s: abs(math.log((parse_ratio(r['ratio']) or 1) / parse_ratio(s))))
            page = next(((p, s) for p in image_pages for s in candidates if s not in p), None)
            if page is None:
                if preferred:
                    slot = preferred
                else:
                    slot = candidates[0]
                new_page = {}
                image_pages.append(new_page)
                page = new_page, slot
            page[0][page[1]] = r
        for p in image_pages:
            plan.append(dict(layout=3, instrument=instrument, slots=p, assets=list(p.values())))
        for r in squares:
            plan.append(dict(layout=4, instrument=instrument, slots={'1:1': r}, assets=[r]))
        for r in videos:
            plan.append(dict(layout=5, instrument=instrument, slots={'video': r}, assets=[r]))
    # Preserve template section ordering: all slide-3 variants, then slide-4, then slide-5.
    return sorted(plan, key=lambda p: p['layout'])


def find_shape(slide, shape_id):
    for s in slide.shapes:
        if s.shape_id == shape_id:
            return s
    raise ValueError(f'Shape ID {shape_id} not found on slide')


def remove_layout_metadata_headers(layout):
    """Remove only the template's exact metadata header text boxes."""
    removed = 0
    for shape in list(layout.shapes):
        if not shape.has_text_frame:
            continue
        if shape.text.strip() in TABLE_HEADER_LABELS:
            element = shape._element
            element.getparent().remove(element)
            removed += 1
    return removed


def description_aligned_geometry(slide, table_shape, layout_name):
    """Use the layout's [Description] placeholder when it exists."""
    description = next((shape for shape in slide.slide_layout.shapes
                        if shape.has_text_frame and '[Description]' in shape.text), None)
    if description is None:
        if layout_name not in TABLE_LAYOUT_WARNINGS:
            print(f'WARNING: {slide.slide_layout.name} has no [Description] placeholder; '
                  'keeping the template table position.')
            TABLE_LAYOUT_WARNINGS.add(layout_name)
        return table_shape.left, table_shape.top

    left = description.left
    minimum_top = description.top + description.height + Cm(0.3)
    top = max(table_shape.top, minimum_top)
    return left, top


def remove_cell_borders(cell):
    """Override table-style borders with transparent edges."""
    tc_pr = cell._tc.get_or_add_tcPr()
    edge_names = ('lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr')
    edge_tags = {f'{{http://schemas.openxmlformats.org/drawingml/2006/main}}{name}'
                 for name in edge_names}
    for child in list(tc_pr):
        if child.tag in edge_tags:
            tc_pr.remove(child)
    for edge_name in edge_names:
        edge = OxmlElement(f'a:{edge_name}')
        edge.append(OxmlElement('a:noFill'))
        tc_pr.append(edge)


def replace_shape(shape, new_shape):
    old = shape._element
    parent = old.getparent()
    index = parent.index(old)
    fresh = new_shape._element
    fresh.getparent().remove(fresh)
    parent.insert(index, fresh)
    parent.remove(old)


def contain_rect(shape, image_path):
    with Image.open(image_path) as im:
        w, h = im.size
    factor = min(shape.width / w, shape.height / h)
    width, height = int(w * factor), int(h * factor)
    return (shape.left + (shape.width-width)//2, shape.top + (shape.height-height)//2, width, height)


def put_image(slide, shape_id, path):
    sh = find_shape(slide, shape_id)
    x, y, w, h = contain_rect(sh, path)
    pic = slide.shapes.add_picture(str(path), x, y, width=w, height=h)
    replace_shape(sh, pic)


def ffmpeg_path():
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except ImportError:
        import shutil
        executable = shutil.which('ffmpeg')
        if executable:
            return executable
        raise RuntimeError('Install imageio-ffmpeg: pip install imageio-ffmpeg')


def make_poster(video, destination):
    # Try the 2-second frame first, then the first frame for short videos.
    # A timeout prevents one problematic file from blocking the entire run.
    for timestamp in ('2', '0'):
        destination.unlink(missing_ok=True)
        command = [
            ffmpeg_path(), '-nostdin', '-hide_banner', '-loglevel', 'error',
            '-y', '-ss', timestamp, '-i', str(video),
            '-frames:v', '1', str(destination)
        ]
        try:
            result = subprocess.run(
                command, capture_output=True, text=True, timeout=60
            )
        except subprocess.TimeoutExpired:
            print(f'WARNING: Thumbnail timed out at {timestamp}s: {video.name}', flush=True)
            continue
        if result.returncode == 0 and destination.is_file() and destination.stat().st_size > 0:
            return
        print(f'WARNING: Thumbnail failed at {timestamp}s: {video.name} '
              f'{result.stderr[-250:]}', flush=True)
    raise RuntimeError(f'Unable to extract thumbnail from: {video}')


def put_video(slide, shape_id, path, temp):
    sh = find_shape(slide, shape_id)
    poster = temp / (asset_id(path.name) + '_poster.jpg')
    make_poster(path, poster)
    # Fit the video inside its placeholder, without stretching or cropping.
    with Image.open(poster) as im:
        width_px, height_px = im.size
    factor = min(sh.width / width_px, sh.height / height_px)
    width, height = int(width_px * factor), int(height_px * factor)
    left = sh.left + (sh.width - width) // 2
    top = sh.top + (sh.height - height) // 2
    movie = slide.shapes.add_movie(str(path), left, top, width, height,
                                   poster_frame_image=str(poster), mime_type='video/mp4')
    replace_shape(sh, movie)


def put_title(slide, instrument):
    sh = find_shape(slide, 2)
    if sh.has_text_frame:
        sh.text = instrument
        for paragraph in sh.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(17 if len(instrument) > 48 else 24)


def put_table(slide, shape_id, assets, layout):
    sh = find_shape(slide, shape_id)
    table_left, table_top = description_aligned_geometry(
        slide, sh, str(layout)
    )
    headers = TABLE_HEADERS if layout in (3, 4) else ['Asset ID', 'Variation', 'Quality', 'Symphony Link']
    data = [headers]
    for r in assets:
        info = r['raw']
        link = str(info.get('Hyperlink', '')).strip()
        if layout in (3, 4):
            data.append([r['id'], str(info.get('ColorSpace', '')).strip() or '[Color Space]',
                         r['ratio'], Path(r['path']).suffix.lstrip('.').upper(), link or '[Symphony Link]'])
        else:
            data.append([r['id'], str(info.get('Perspective', '')).strip() or '[Variation]',
                         '[Quality]', link or '[Symphony Link]'])

    # URLs remain in the source data for hyperlink targets, while their visible
    # label is short and consistent in both sizing and rendering.
    display_data = [list(row) for row in data]
    for ri in range(1, len(data)):
        link_value = str(data[ri][-1])
        if link_value.startswith('https://'):
            display_data[ri][-1] = 'Symphony Link'

    font_size = 8 if layout == 3 else 9
    margins_pt = 2.5
    wrap_columns = (1,) if layout == 5 else ()
    widths, font_size = _table_column_widths(
        display_data, sh.width, font_size, margins_pt, wrap_columns
    )
    frame = slide.shapes.add_table(len(data), len(headers), table_left, table_top, sh.width, sh.height)
    replace_shape(sh, frame)
    table = frame.table
    for col, width in zip(table.columns, widths):
        col.width = width

    vertical_margin_pt = 1.5
    row_heights = [Cm(1)]
    for ri, row in enumerate(display_data):
        max_lines = 1
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = str(val)
            cell.fill.background()
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Pt(margins_pt)
            cell.margin_top = cell.margin_bottom = Pt(vertical_margin_pt)
            cell.text_frame.word_wrap = ci in wrap_columns
            remove_cell_borders(cell)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT
                paragraph.space_before = Pt(0)
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = TABLE_FONT
                    run.font.size = Pt(font_size)
                    run.font.bold = ri == 0
                    if ri == 0:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                    if (ri > 0 and ci == len(headers) - 1
                            and str(data[ri][ci]).startswith('https://')):
                        run.hyperlink.address = str(data[ri][ci])

            if ri > 0:
                content_width_pt = widths[ci] / 12700 - 2 * margins_pt
                lines = (_wrapped_line_count(val, content_width_pt, font_size)
                         if ci in wrap_columns else 1)
                max_lines = max(max_lines, lines)
        if ri > 0:
            # 1.2 line-height plus the small vertical margins gives compact rows;
            # only the Variation column can increase a row beyond one line.
            content_height = Pt(font_size * 1.2 * max_lines + 2 * vertical_margin_pt)
            row_heights.append(max(Cm(0.8), content_height))

    for row_index, height in enumerate(row_heights):
        table.rows[row_index].height = height
    # The original placeholder is deliberately tall; shrink the table to its
    # actual content height so a single asset sits immediately below the header.
    frame.height = sum(row_heights)


def generate(template, plan, output):
    prs = Presentation(template)
    # Reuse slide layouts so every new slide retains the branded masters and placeholders.
    # Capture placeholder indices: new slides have different shape IDs from the originals.
    original_placeholder_idx = {number: {sh.shape_id: sh.placeholder_format.idx
                                  for sh in prs.slides[number-1].shapes if sh.is_placeholder}
                                for number in (3, 4, 5)}
    # Remove the three sample content slides; retain cover and Slide 2.
    slide_ids = prs.slides._sldIdLst
    for index in (4, 3, 2):
        node = slide_ids[index]
        prs.part.drop_rel(node.rId)
        slide_ids.remove(node)
    # Locate layouts by names from inspected master template.
    names = {3: 'KV_Crops_V1', 4: 'KV Master only 1:1', 5: 'Product Film'}
    resolved = {number: next((l for l in prs.slide_layouts if l.name == name), None)
                for number, name in names.items()}
    if any(v is None for v in resolved.values()):
        raise ValueError('Required template slide layouts not found: ' + str(names))
    # The duplicate labels are text boxes on these layouts. Remove only those
    # exact labels in this in-memory output presentation; the source PPTX is read-only.
    for layout in resolved.values():
        retained_users = [s for s in prs.slides if s.slide_layout == layout]
        if retained_users:
            raise ValueError(
                f'Cannot safely remove metadata headers from layout {layout.name}; '
                'a retained template slide still uses it.'
            )
        removed = remove_layout_metadata_headers(layout)
        if not removed:
            print(f'WARNING: No template metadata header text boxes found on {layout.name}')
        else:
            print(f'Removed {removed} inherited metadata headers from {layout.name}')
    with tempfile.TemporaryDirectory() as td:
        temp = Path(td)
        for index, page in enumerate(plan, 1):
            slide = prs.slides.add_slide(resolved[page['layout']])
            def sid(original_id):
                idx = original_placeholder_idx[page['layout']][original_id]
                return next(sh.shape_id for sh in slide.shapes
                            if sh.is_placeholder and sh.placeholder_format.idx == idx)
            put_title(slide, page['instrument'])
            if page['layout'] == 3:
                for ratio, r in page['slots'].items():
                    put_image(slide, sid(IMAGE_SLOTS[ratio]), r['path'])
                put_table(slide, sid(8), page['assets'], 3)
            elif page['layout'] == 4:
                put_image(slide, sid(3), page['assets'][0]['path'])
                put_table(slide, sid(8), page['assets'], 4)
            else:
                video_path = page['assets'][0]['path']
                print(f'Processing video: {video_path.name}', flush=True)
                put_video(slide, sid(6), video_path, temp)
                put_table(slide, sid(7), page['assets'], 5)
            print(f'[{index}/{len(plan)}] {page["instrument"]}: {page["layout"]} ({len(page["assets"])} assets)', flush=True)
    prs.save(output)
    print('Saved:', output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--plan-only', action='store_true', help='Validate Excel join and plan without media files')
    args = parser.parse_args()
    root = project_root()
    inp = root / 'input'
    excel = inp / 'assets_table1_cleaned.xlsx'
    raw = inp / 'Export_report_20260917_113237_Raw(1).xlsx'
    template = inp / 'master_template.pptx'
    for p in (excel, raw, template):
        if not p.is_file():
            raise FileNotFoundError(f'Missing input: {p}')
    records = read_assets(excel, raw)
    if args.plan_only:
        for r in records:
            r['path'] = Path(r['name'])
    else:
        find_local_files(records, inp / 'Assets')
    missing = [r for r in records if not r.get('path')]
    if missing:
        print(f'WARNING: {len(missing)} assets could not be located unambiguously')
        for r in missing:
            print(f'  Excel row {r["index"]}: {r["name"]} ({r.get("problem")})')
    plan = group_and_allocate(records)
    manifest = root / 'automation_manifest.csv'
    with manifest.open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(['output_slide', 'template_slide', 'instrument', 'asset_id', 'file_name', 'ratio', 'local_path'])
        for slide_no, p in enumerate(plan, 3):
            for r in p['assets']:
                writer.writerow([slide_no, p['layout'], p['instrument'], r['id'], r['name'], r['ratio'], r['path']])
    print(f'Excel assets: {len(records)} | Found: {len(records)-len(missing)} | Planned slides: {len(plan)}')
    print('Manifest:', manifest)
    if not args.plan_only:
        generate(template, plan, root / 'final_output.pptx')


if __name__ == '__main__':
    main()
