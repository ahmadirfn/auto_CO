from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
import pandas as pd
import re


# ==================================================
# 1. PATH CONFIGURATION
# ==================================================

BASE_DIR = Path(__file__).resolve().parent.parent

INPUT_DIR = BASE_DIR / "input"

PPTX_PATH = INPUT_DIR / "master_template.pptx"
EXCEL_PATH = INPUT_DIR / "assets_table1.xlsx"
OUTPUT_PATH = BASE_DIR / "output_table.pptx"

# Slide indices are zero-based
TABLE_SLIDE_INDEX = 2
TABLE_SHAPE_ID = 8


# ==================================================
# 2. READ TABLE DATA FROM EXCEL
# ==================================================

df = pd.read_excel(EXCEL_PATH, header=2)

df.columns = df.columns.astype(str).str.strip().str.lower()

if "file_name" not in df.columns:
    raise ValueError("Column 'file_name' not found in Excel.")

df = df.dropna(subset=["file_name"])


def get_asset_id(filename):
    match = re.match(r"(\d+)", str(filename).strip())
    return match.group(1) if match else ""


TABLE_DATA = [
    ["Asset ID", "Variation", "Quality", "Symphony Link"]
]

for _, row in df.iterrows():
    TABLE_DATA.append([
        get_asset_id(row["file_name"]),
        "",
        "",
        ""
    ])

print(f"Excel records found: {len(df)}")


# ==================================================
# 3. FIND SHAPE BY ID
# ==================================================

def find_shape(slide, shape_id):

    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape

    raise ValueError(f"Shape ID {shape_id} not found.")


# ==================================================
# 4. REPLACE ORIGINAL SHAPE
# ==================================================

def replace_shape(original_shape, new_shape):

    old_element = original_shape._element
    new_element = new_shape._element

    parent = old_element.getparent()
    index = parent.index(old_element)

    new_element.getparent().remove(new_element)
    parent.insert(index, new_element)

    parent.remove(old_element)


# ==================================================
# 5. INSERT TABLE
# ==================================================

def insert_table(slide, shape_id, data):

    if not data or not data[0]:
        raise ValueError("Table data cannot be empty.")

    rows = len(data)
    cols = len(data[0])

    shape = find_shape(slide, shape_id)

    print(f"Table placeholder found: {shape.name}")

    # Native table placeholder
    if (
        shape.is_placeholder
        and shape.placeholder_format.type == PP_PLACEHOLDER.TABLE
    ):
        graphic_frame = shape.insert_table(
            rows=rows,
            cols=cols
        )

        table = graphic_frame.table

    else:
        graphic_frame = slide.shapes.add_table(
            rows=rows,
            cols=cols,
            left=shape.left,
            top=shape.top,
            width=shape.width,
            height=shape.height
        )

        table = graphic_frame.table
        replace_shape(shape, graphic_frame)

    # --------------------------------------------------
    # POPULATE TABLE
    # --------------------------------------------------

    for row_idx, row_data in enumerate(data):
        for col_idx, value in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(value) if value is not None else ""

    # --------------------------------------------------
    # FORMAT TABLE
    # --------------------------------------------------

    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:

            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            cell.margin_left = Pt(5)
            cell.margin_right = Pt(5)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)

            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER

                for run in paragraph.runs:
                    run.font.size = Pt(10)

                    if row_idx == 0:
                        run.font.bold = True

    # --------------------------------------------------
    # SET COLUMN WIDTHS
    # --------------------------------------------------

    if cols == 4:
        column_ratios = [0.20, 0.35, 0.15, 0.30]

        total_width = sum(
            col.width for col in table.columns
        )

        widths = [
            int(total_width * ratio)
            for ratio in column_ratios[:-1]
        ]

        widths.append(total_width - sum(widths))

        for col, width in zip(table.columns, widths):
            col.width = width

    print(f"Table inserted: {rows} rows x {cols} columns")

    return table


# ==================================================
# 6. MAIN AUTOMATION
# ==================================================

def main():

    if not PPTX_PATH.is_file():
        raise FileNotFoundError(
            f"PowerPoint template not found: {PPTX_PATH}"
        )

    prs = Presentation(PPTX_PATH)

    if TABLE_SLIDE_INDEX >= len(prs.slides):
        raise IndexError("Configured slide does not exist.")

    insert_table(
        slide=prs.slides[TABLE_SLIDE_INDEX],
        shape_id=TABLE_SHAPE_ID,
        data=TABLE_DATA
    )

    prs.save(str(OUTPUT_PATH))

    print(f"\nPowerPoint saved successfully: {OUTPUT_PATH}")


# ==================================================
# 7. EXECUTE
# ==================================================

if __name__ == "__main__":
    main()