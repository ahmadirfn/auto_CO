from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pathlib import Path
from PIL import Image
import pandas as pd
import subprocess
import tempfile
import shutil

# ==================================================
# 1. PATH CONFIGURATION
# ==================================================

BASE_DIR = Path(__file__).resolve().parent.parent

INPUT_DIR = BASE_DIR / "input"
ASSETS_DIR = INPUT_DIR / "Assets"

PPTX_PATH = INPUT_DIR / "master_template.pptx"
EXCEL_PATH = INPUT_DIR / "assets_table1.xlsx"
OUTPUT_PATH = BASE_DIR / "output.pptx"

# Thumbnail timestamp in seconds
THUMBNAIL_TIMESTAMP = 2


# ==================================================
# 2. READ ASSETS FROM EXCEL
# ==================================================

df = pd.read_excel(EXCEL_PATH, header=2)

df.columns = df.columns.astype(str).str.strip().str.lower()
df["format"] = df["format"].astype(str).str.strip().str.lower()

print("Excel columns:", df.columns.tolist())

image_rows = df[df["format"].isin(["image", "jpg", "jpeg", "png"])]
video_rows = df[df["format"].isin(["video", "mp4", "mov"])]


# Search for actual files inside input/Assets
def find_asset(filename):

    filename = str(filename).strip()

    matches = [
        p for p in ASSETS_DIR.rglob("*")
        if p.is_file()
        and (
            p.name.lower() == filename.lower()
            or p.stem.lower() == filename.lower()
        )
    ]

    if not matches:
        raise FileNotFoundError(
            f"Asset not found in {ASSETS_DIR}: {filename}"
        )

    return matches[0]


IMAGE_PATH = None
VIDEO_PATH = None

if not image_rows.empty:
    IMAGE_PATH = find_asset(image_rows.iloc[0]["file_name"])

if not video_rows.empty:
    VIDEO_PATH = find_asset(video_rows.iloc[0]["file_name"])


# Validate files
for path in [PPTX_PATH, EXCEL_PATH, IMAGE_PATH, VIDEO_PATH]:
    if path is not None and not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

if VIDEO_PATH is not None and shutil.which("ffmpeg") is None:
    raise RuntimeError(
        "FFmpeg not found. Install FFmpeg and add it to PATH."
    )

print(f"Selected image: {IMAGE_PATH}")
print(f"Selected video: {VIDEO_PATH}")


# ==================================================
# 3. HELPER: FIND SHAPE BY ID
# ==================================================

def find_shape(slide, shape_id):

    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape

    raise ValueError(f"Shape ID {shape_id} not found")


# ==================================================
# 4. HELPER: REPLACE SHAPE
# ==================================================

def replace_shape(original_shape, new_shape):

    old_element = original_shape._element
    new_element = new_shape._element

    parent = old_element.getparent()
    index = parent.index(old_element)

    new_element.getparent().remove(new_element)
    parent.insert(index, new_element)
    parent.remove(old_element)


# ==================================================
# 5. HELPER: INSERT IMAGE
# ==================================================

def insert_image(slide, shape_id, image_path):

    shape = find_shape(slide, shape_id)

    print(f"Found image placeholder: {shape.name}")

    if (
        shape.is_placeholder
        and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ):
        shape.insert_picture(str(image_path))
        print("Image inserted into picture placeholder.")
        return

    picture = slide.shapes.add_picture(
        str(image_path),
        shape.left,
        shape.top,
        width=shape.width,
        height=shape.height
    )

    replace_shape(shape, picture)

    print("Image replaced existing placeholder.")


# ==================================================
# 6. EXTRACT VIDEO THUMBNAIL
# ==================================================

def extract_video_thumbnail(video_path, output_path, timestamp=2):

    print(f"\nExtracting thumbnail at {timestamp}s...")

    command = [
        "ffmpeg",
        "-y",
        "-ss", str(timestamp),
        "-i", str(video_path),
        "-frames:v", "1",
        "-q:v", "2",
        str(output_path)
    ]

    result = subprocess.run(
        command,
        capture_output=True,
        text=True
    )

    if result.returncode != 0:
        raise RuntimeError(
            f"FFmpeg failed:\n{result.stderr}"
        )

    if not output_path.is_file():
        raise RuntimeError("Thumbnail was not generated.")

    with Image.open(output_path) as img:
        img.verify()

    print(f"Thumbnail generated: {output_path}")

    return output_path


# ==================================================
# 7. HELPER: INSERT VIDEO WITH THUMBNAIL
# ==================================================

def insert_video(slide, shape_id, video_path, thumbnail_path):

    shape = find_shape(slide, shape_id)

    print(f"Found video placeholder: {shape.name}")

    movie = slide.shapes.add_movie(
        movie_file=str(video_path),
        left=shape.left,
        top=shape.top,
        width=shape.width,
        height=shape.height,
        poster_frame_image=str(thumbnail_path),
        mime_type="video/mp4"
    )

    replace_shape(shape, movie)

    print("Video embedded with thumbnail successfully.")


# ==================================================
# 8. INSERT ASSETS
# ==================================================

prs = Presentation(PPTX_PATH)

# Slide 2 - Image
if IMAGE_PATH is not None:

    insert_image(
        slide=prs.slides[1],
        shape_id=10,
        image_path=IMAGE_PATH
    )

else:
    print("No image found in Excel. Skipping image insertion.")


# Slide 5 - Video
if VIDEO_PATH is not None:

    with tempfile.TemporaryDirectory() as temp_dir:

        thumbnail_path = (
            Path(temp_dir) / "video_thumbnail.jpg"
        )

        extract_video_thumbnail(
            video_path=VIDEO_PATH,
            output_path=thumbnail_path,
            timestamp=THUMBNAIL_TIMESTAMP
        )

        insert_video(
            slide=prs.slides[4],
            shape_id=6,
            video_path=VIDEO_PATH,
            thumbnail_path=thumbnail_path
        )

        # Save while thumbnail exists
        prs.save(str(OUTPUT_PATH))

else:
    print("No video found in Excel. Skipping video insertion.")
    prs.save(str(OUTPUT_PATH))


# ==================================================
# 9. SAVE PRESENTATION
# ==================================================

print(f"\nPowerPoint saved successfully: {OUTPUT_PATH}")