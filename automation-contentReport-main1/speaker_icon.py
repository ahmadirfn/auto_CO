from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pathlib import Path
import pandas as pd

# ==================================================
# 1. PATH CONFIGURATION
# ==================================================

BASE_DIR = Path(__file__).resolve().parent.parent

INPUT_DIR = BASE_DIR / "input"
ASSETS_DIR = INPUT_DIR / "Assets"

PPTX_PATH = INPUT_DIR / "master_template.pptx"
EXCEL_PATH = INPUT_DIR / "assets_table1.xlsx"
OUTPUT_PATH = BASE_DIR / "output.pptx"


# ==================================================
# 2. READ ASSETS FROM EXCEL
# ==================================================

df = pd.read_excel(EXCEL_PATH, header=2)

df.columns = df.columns.astype(str).str.strip().str.lower()
df["format"] = df["format"].astype(str).str.strip().str.lower()

print("Excel columns:", df.columns.tolist())

image_rows = df[df["format"].isin(["image", "jpg", "jpeg", "png"])]
video_rows = df[df["format"].isin(["video", "mp4", "mov"])]

# Search for actual files inside input/Assets
def find_asset(filename):
    filename = str(filename).strip()

    matches = [
        p for p in ASSETS_DIR.rglob("*")
        if p.is_file()
        and (
            p.name.lower() == filename.lower()
            or p.stem.lower() == filename.lower()
        )
    ]

    if not matches:
        raise FileNotFoundError(
            f"Asset not found in {ASSETS_DIR}: {filename}"
        )

    return matches[0]


IMAGE_PATH = None
VIDEO_PATH = None

if not image_rows.empty:
    IMAGE_PATH = find_asset(image_rows.iloc[0]["file_name"])

if not video_rows.empty:
    VIDEO_PATH = find_asset(video_rows.iloc[0]["file_name"])

# Validate paths
for path in [PPTX_PATH, EXCEL_PATH, IMAGE_PATH, VIDEO_PATH]:
    if path is not None and not path.is_file():
        raise FileNotFoundError(f"File not found: {path}")

print(f"Selected image: {IMAGE_PATH}")
print(f"Selected video: {VIDEO_PATH}")

# Open existing PowerPoint
prs = Presentation(PPTX_PATH)


# ==================================================
# 3. HELPER: FIND SHAPE BY ID
# ==================================================

def find_shape(slide, shape_id):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape

    raise ValueError(f"Shape ID {shape_id} not found")


# ==================================================
# 4. HELPER: REPLACE SHAPE WITH IMAGE
# ==================================================

def insert_image(slide, shape_id, image_path):

    shape = find_shape(slide, shape_id)

    print(f"Found image placeholder: {shape.name}")

    # Use native picture placeholder if supported
    if (
        shape.is_placeholder
        and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    ):
        shape.insert_picture(str(image_path))
        print("Image inserted into picture placeholder.")
        return

    # Otherwise replace the existing shape
    left = shape.left
    top = shape.top
    width = shape.width
    height = shape.height

    old_element = shape._element
    parent = old_element.getparent()
    index = parent.index(old_element)

    picture = slide.shapes.add_picture(
        str(image_path),
        left,
        top,
        width=width,
        height=height
    )

    # Preserve stacking position
    parent.remove(picture._element)
    parent.insert(index, picture._element)

    # Remove original placeholder
    parent.remove(old_element)

    print("Image replaced existing placeholder.")


# ==================================================
# 5. HELPER: REPLACE SHAPE WITH VIDEO
# ==================================================

def insert_video(slide, shape_id, video_path):

    shape = find_shape(slide, shape_id)

    print(f"Found video placeholder: {shape.name}")

    left = shape.left
    top = shape.top
    width = shape.width
    height = shape.height

    old_element = shape._element
    parent = old_element.getparent()
    index = parent.index(old_element)

    movie = slide.shapes.add_movie(
        movie_file=str(video_path),
        left=left,
        top=top,
        width=width,
        height=height,
        mime_type="video/mp4"
    )

    # Preserve stacking position
    parent.remove(movie._element)
    parent.insert(index, movie._element)

    # Remove original placeholder
    parent.remove(old_element)

    print("Video embedded successfully.")


# ==================================================
# 6. INSERT ASSETS
# ==================================================

# Slide 2 - Image
if IMAGE_PATH is not None:
    insert_image(
        slide=prs.slides[1],
        shape_id=10,
        image_path=IMAGE_PATH
    )
else:
    print("No image found in Excel. Skipping image insertion.")

# Slide 5 - Video
if VIDEO_PATH is not None:
    insert_video(
        slide=prs.slides[4],
        shape_id=6,
        video_path=VIDEO_PATH
    )
else:
    print("No video found in Excel. Skipping video insertion.")


# ==================================================
# 7. SAVE PRESENTATION
# ==================================================

prs.save(str(OUTPUT_PATH))

print(f"\nPowerPoint saved successfully: {OUTPUT_PATH}")